from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from app.main import app


def test_export_preserves_reviewed_text_and_direction():
    text = "اعلامیہ\n\nنام: [FULL NAME]\nCNIC: 12345-1234567-1\nمنظور نہیں ہوئی۔"
    with TestClient(app) as client:
        result = client.post('/api/export/docx', json={"text":text,"direction":"en-ur","approved":True})
    assert result.status_code == 200
    doc = Document(BytesIO(result.content))
    assert '\n'.join(p.text for p in doc.paragraphs) == text
    with ZipFile(BytesIO(result.content)) as archive:
        xml = archive.read('word/document.xml').decode()
        assert 'w:bidi w:val="1"' in xml
        assert 'w:rtl w:val="0"' in xml
        assert not any('vba' in name for name in archive.namelist())
    assert result.headers['cache-control'] == 'no-store'


def test_unapproved_empty_and_invalid_exports_rejected():
    with TestClient(app) as client:
        for extra in ({"approved":False}, {"text":" "}, {"text":"bad\x00text"}, {"text":"x"*30001}, {"direction":"fr-en"}):
            body = {"text":"Reviewed text", "direction":"ur-en", "approved":True, **extra}
            assert client.post('/api/export/docx', json=body).status_code == 422
        assert client.post('/api/export/docx', json={"text":"Text", "direction":"ur-en"}).status_code == 422


def test_english_export_and_request_size():
    with TestClient(app) as client:
        result = client.post('/api/export/docx', json={"text":"Not approved.\nAmount: 25,000", "direction":"ur-en", "approved":True})
        assert result.status_code == 200
        assert Document(BytesIO(result.content)).paragraphs[0].text == 'Not approved.'
        assert client.post('/api/export/docx', content=b'x'*262145).status_code == 413
