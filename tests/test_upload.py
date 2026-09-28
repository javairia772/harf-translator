from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from docx import Document
from fastapi.testclient import TestClient
from app.main import app, get_provider
from app.export import word_document
from app.translation import ModelTranslation

MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def upload(client, data):
    return client.post('/api/upload', content=data, headers={'content-type':MIME})


@pytest.mark.parametrize('direction,source,target', [
    ('en-ur','Leave has not been approved.','رخصت منظور نہیں ہوئی۔'),
    ('en-ur','I, [FULL NAME], shall repay Rs. 25,000 within 30 days.','میں، [FULL NAME]، 30 دنوں کے اندر 25,000 روپے واپس ادا کروں گا۔'),
    ('ur-en','درخواست 15 اکتوبر 2026 تک جمع کروائیں۔','Submit the application by 15 October 2026.'),
    ('ur-en','میں تصدیق کرتا ہوں کہ کوئی اہم حقیقت پوشیدہ نہیں رکھی گئی۔','I confirm that no material fact has been concealed.'),
])
def test_upload_translate_approve_export_both_directions(direction, source, target):
    class FixtureProvider:
        async def translate(self, body, progress=None):
            assert body.text == source and body.direction == direction
            return ModelTranslation(translation=target, notes=[])
    app.dependency_overrides[get_provider] = FixtureProvider
    try:
        with TestClient(app) as client:
            doc = word_document(source, 'ur-en' if direction == 'en-ur' else 'en-ur')
            extracted = upload(client, doc)
            assert extracted.status_code == 200
            assert extracted.json()['text'] == source
            response = client.post('/api/translate', json={'text':extracted.json()['text'], 'direction':direction})
            assert response.status_code == 200
            result = client.post('/api/export/docx', json={'text':response.json()['translation'], 'direction':direction, 'approved':True})
            assert result.status_code == 200
            assert '\n'.join(p.text for p in Document(BytesIO(result.content)).paragraphs) == target
    finally:
        app.dependency_overrides.clear()


def test_table_order_and_blank_paragraphs():
    doc = Document(); doc.add_paragraph('Before'); doc.add_paragraph('')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0,0).text = 'Name'; table.cell(0,1).text = '[FULL NAME]'
    doc.add_paragraph('After'); buffer = BytesIO(); doc.save(buffer)
    with TestClient(app) as client:
        result = upload(client, buffer.getvalue())
    assert result.json()['text'] == 'Before\n\nName\t[FULL NAME]\nAfter'
    assert len(result.json()['warnings']) == 2


def test_reject_corrupt_oversized_empty_and_long_documents():
    with TestClient(app) as client:
        assert upload(client,b'not a docx').status_code == 422
        assert upload(client,b'x'*(2*1024*1024+1)).status_code == 413
        assert upload(client,word_document(' ', 'ur-en')).status_code == 422
        assert upload(client,word_document('x'*5001, 'ur-en')).status_code == 422
        assert client.post('/api/upload',content=b'%PDF',headers={'content-type':'application/pdf'}).status_code == 422
        assert client.post('/api/upload',content=b'old doc',headers={'content-type':'application/msword'}).status_code == 415


def test_header_and_revision_rejected_without_silent_omissions():
    doc = Document(); doc.add_paragraph('Body'); doc.sections[0].header.paragraphs[0].text = 'Important condition'
    buf = BytesIO(); doc.save(buf)
    with TestClient(app) as client:
        assert upload(client,buf.getvalue()).status_code == 422
        src = word_document('Body', 'ur-en'); out=BytesIO()
        with ZipFile(BytesIO(src)) as original, ZipFile(out,'w',ZIP_DEFLATED) as changed:
            for item in original.infolist():
                data=original.read(item.filename)
                if item.filename == 'word/document.xml':
                    data=data.replace(b'<w:t>Body</w:t>',b'<w:del><w:r><w:delText>Hidden</w:delText></w:r></w:del>')
                changed.writestr(item.filename,data)
        assert upload(client,out.getvalue()).status_code == 422


def test_zip_expansion_rejected():
    out=BytesIO()
    with ZipFile(out,'w',ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml',b'x'*(8*1024*1024+1))
        archive.writestr('[Content_Types].xml',b'x')
    with TestClient(app) as client:
        assert upload(client,out.getvalue()).status_code == 422
