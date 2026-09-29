"""Export reviewed plain text without a model call or stored document."""
from io import BytesIO
from typing import Literal
from uuid import UUID
import re

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pydantic import BaseModel, ConfigDict, Field, field_validator


class ExportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    text: str = Field(min_length=1, max_length=30000)
    direction: Literal["en-ur", "ur-en"]
    approved: Literal[True]
    workflow_id: UUID | None = None

    @field_validator("text")
    @classmethod
    def valid_text(cls, value):
        if not value.strip() or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", value):
            raise ValueError("Text contains unsupported characters or is empty")
        return value


def word_document(text: str, direction: str) -> bytes:
    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(0)
    urdu = direction == "en-ur"
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        paragraph = document.add_paragraph()
        properties = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1" if urdu else "0")
        properties.append(bidi)
        alignment = OxmlElement("w:jc")
        alignment.set(qn("w:val"), "right" if urdu else "left")
        properties.append(alignment)
        # Separate Arabic-script and Latin runs so CNIC/placeholder order stays readable.
        for part in re.findall(r"[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff]+|[^\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff]+", line):
            run = paragraph.add_run(part)
            rtl = bool(re.match(r"[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff]", part))
            run.font.rtl = rtl
            props = run._r.get_or_add_rPr()
            fonts = props.get_or_add_rFonts()
            fonts.set(qn("w:cs"), "Arial")
            size = OxmlElement("w:szCs")
            size.set(qn("w:val"), "24")
            props.append(size)
            language = OxmlElement("w:lang")
            language.set(qn("w:val"), "ur-PK" if rtl else "en-GB")
            language.set(qn("w:bidi"), "ur-PK")
            props.append(language)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
